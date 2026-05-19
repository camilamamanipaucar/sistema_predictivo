import io
import re
import warnings
import numpy as np
import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt
import seaborn as sns


from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation

from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from sklearn.model_selection import train_test_split
from sklearn.metrics import (
    mean_absolute_error,
    mean_squared_error,
    r2_score,
    confusion_matrix,
    accuracy_score,
    precision_score,
    recall_score,
    f1_score
)
from sklearn.linear_model import LinearRegression
from sklearn.pipeline import Pipeline
from sklearn.compose import ColumnTransformer
from sklearn.preprocessing import OneHotEncoder

import statsmodels.api as sm
from scipy import stats
from statsmodels.stats.outliers_influence import variance_inflation_factor

warnings.filterwarnings("ignore")

# ==========================================================
# CONFIGURACIÓN GENERAL
# ==========================================================

st.set_page_config(
    page_title="Sistema Predictivo Inteligente",
    page_icon="📊",
    layout="wide"
)

st.markdown(
    """
    <style>
    .main {
        animation: fadeIn 0.8s;
    }
    @keyframes fadeIn {
        from {opacity: 0;}
        to {opacity: 1;}
    }
    .titulo {
        font-size: 34px;
        font-weight: bold;
        color: #1f4e79;
    }
    .subtitulo {
        font-size: 18px;
        color: #444;
    }
    </style>
    """,
    unsafe_allow_html=True
)

st.markdown('<div class="titulo">📊 Sistema Web Inteligente de Análisis Predictivo</div>', unsafe_allow_html=True)
st.markdown('<div class="subtitulo">Análisis de Excel, series de tiempo, regresión múltiple, matriz de confusión y documentos del docente</div>', unsafe_allow_html=True)
st.markdown("---")

# ==========================================================
# FUNCIONES GENERALES
# ==========================================================

def limpiar_dataframe(df):
    df = df.copy()
    df = df.dropna(axis=1, how="all")
    columnas_eliminar = [col for col in df.columns if str(col).lower().startswith("unnamed")]
    df = df.drop(columns=columnas_eliminar, errors="ignore")
    return df


def cargar_archivo_excel(archivo):
    nombre = archivo.name.lower()

    # Lee CSV aunque el nombre venga raro, por ejemplo: student-mat.csv_e89m3c7z7h
    if ".csv" in nombre:
        try:
            archivo.seek(0)
            return pd.read_csv(archivo, sep=None, engine="python")
        except UnicodeDecodeError:
            archivo.seek(0)
            return pd.read_csv(archivo, sep=None, engine="python", encoding="latin1")

    elif nombre.endswith(".xlsx") or nombre.endswith(".xls"):
        return pd.read_excel(archivo)

    else:
        try:
            archivo.seek(0)
            return pd.read_csv(archivo, sep=None, engine="python")
        except Exception:
            st.error("El archivo no pudo leerse. Renómbralo como .csv o .xlsx.")
            st.stop()


def obtener_columnas_numericas(df):
    return df.select_dtypes(include=[np.number]).columns.tolist()


def obtener_columnas_categoricas(df):
    return df.select_dtypes(include=["object", "category", "bool"]).columns.tolist()


def mape_seguro(y_true, y_pred):
    y_true = np.array(y_true)
    y_pred = np.array(y_pred)
    mascara = y_true != 0
    if mascara.sum() == 0:
        return np.nan
    return np.mean(np.abs((y_true[mascara] - y_pred[mascara]) / y_true[mascara])) * 100


def calcular_metricas_regresion(y_true, y_pred):
    mae = mean_absolute_error(y_true, y_pred)
    mse = mean_squared_error(y_true, y_pred)
    rmse = np.sqrt(mse)
    r2 = r2_score(y_true, y_pred)
    return mae, mse, rmse, r2


def crear_pdf_reporte(variable_objetivo, r2, mae, accuracy):
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)

    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, 750, "Reporte del Sistema Predictivo Inteligente")

    c.setFont("Helvetica", 12)
    c.drawString(50, 710, f"Variable objetivo: {variable_objetivo}")
    c.drawString(50, 685, f"R² del modelo múltiple: {r2:.4f}" if r2 is not None else "R² del modelo múltiple: No calculado")
    c.drawString(50, 660, f"MAE: {mae:.4f}" if mae is not None else "MAE: No calculado")
    c.drawString(50, 635, f"Accuracy de clasificación: {accuracy:.4f}" if accuracy is not None else "Accuracy de clasificación: No calculado")

    c.drawString(50, 590, "Interpretación:")
    c.drawString(50, 565, "El reporte resume los principales resultados obtenidos del análisis.")
    c.drawString(50, 540, "El mejor modelo se elige según el menor error o mejor rendimiento.")

    c.save()
    buffer.seek(0)
    return buffer


# ==========================================================
# MÉTODOS DE PRONÓSTICO
# ==========================================================

def metodo_ingenuo(train, horizonte):
    return np.repeat(train[-1], horizonte)


def metodo_media(train, horizonte):
    return np.repeat(np.mean(train), horizonte)


def metodo_media_movil(train, horizonte, ventana):
    ventana = max(1, min(ventana, len(train)))
    valores = list(train[-ventana:])
    pronosticos = []

    for _ in range(horizonte):
        pred = np.mean(valores[-ventana:])
        pronosticos.append(pred)
        valores.append(pred)

    return np.array(pronosticos)


def metodo_deriva(train, horizonte):
    if len(train) < 2:
        return np.repeat(train[-1], horizonte)

    pendiente = (train[-1] - train[0]) / (len(train) - 1)
    return np.array([train[-1] + pendiente * (i + 1) for i in range(horizonte)])


def metodo_ingenuo_estacional(train, horizonte, temporada):
    temporada = max(1, min(temporada, len(train)))
    ultimos = train[-temporada:]
    pronosticos = []

    for i in range(horizonte):
        pronosticos.append(ultimos[i % temporada])

    return np.array(pronosticos)


# ==========================================================
# DOCUMENTOS
# ==========================================================

def leer_pdf(archivo):
    texto = ""
    reader = PdfReader(archivo)
    for i, page in enumerate(reader.pages):
        contenido = page.extract_text()
        if contenido:
            texto += f"\n\n--- Página {i+1} ---\n"
            texto += contenido
    return texto


def leer_docx(archivo):
    doc = Document(archivo)
    texto = ""
    for parrafo in doc.paragraphs:
        texto += parrafo.text + "\n"
    return texto


def leer_pptx(archivo):
    prs = Presentation(archivo)
    texto = ""
    for i, slide in enumerate(prs.slides):
        texto += f"\n\n--- Diapositiva {i+1} ---\n"
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                texto += shape.text + "\n"
    return texto


def detectar_tema(texto):
    texto_min = texto.lower()

    temas = []

    if "serie" in texto_min or "pronóstico" in texto_min or "pronostico" in texto_min:
        temas.append("Series de tiempo y pronósticos")
    if "regresión" in texto_min or "regresion" in texto_min:
        temas.append("Regresión lineal")
    if "matriz de confusión" in texto_min or "confusion" in texto_min:
        temas.append("Matriz de confusión")
    if "clasificación" in texto_min or "clasificacion" in texto_min:
        temas.append("Clasificación")
    if "anova" in texto_min:
        temas.append("ANOVA")
    if "coeficiente" in texto_min:
        temas.append("Coeficientes del modelo")

    if not temas:
        temas.append("Tema no identificado automáticamente")

    return temas


def extraer_posibles_formulas(texto):
    patrones = [
        r"y\s*=\s*a\s*\+\s*b",
        r"Y\s*=\s*a\s*\+.*",
        r"R2|R²",
        r"MAE|MSE|RMSE|MAPE",
        r"VP|VN|FP|FN",
        r"accuracy|precision|recall",
    ]

    encontrados = []
    for patron in patrones:
        resultados = re.findall(patron, texto, flags=re.IGNORECASE)
        encontrados.extend(resultados)

    return list(set(encontrados))


# ==========================================================
# MENÚ PRINCIPAL
# ==========================================================

st.sidebar.title("📌 Menú principal")

opcion_principal = st.sidebar.radio(
    "Selecciona una funcionalidad",
    [
        "📁 Análisis de Excel",
        "📄 Documentos del Inge"
    ]
)

# ==========================================================
# FUNCIONALIDAD 1: EXCEL
# ==========================================================

if opcion_principal == "📁 Análisis de Excel":

    st.header("📁 Funcionalidad 1: Análisis de Excel")

    archivo = st.sidebar.file_uploader(
        "Sube un archivo CSV o Excel",
        type=None
    )

    df = None

    if archivo is not None:
        try:
            df = cargar_archivo_excel(archivo)
            df = limpiar_dataframe(df)
            st.sidebar.success("Archivo cargado correctamente")
        except Exception as e:
            st.error(f"Error al cargar archivo: {e}")
            st.stop()
    else:
        try:
            df = pd.read_csv("ds_salaries.csv")
            df = limpiar_dataframe(df)
            st.sidebar.info("Se usó el archivo ds_salaries.csv por defecto")
        except:
            st.info("Sube un archivo CSV o Excel para iniciar.")
            st.stop()

    if df is None or df.empty:
        st.warning("No hay datos disponibles.")
        st.stop()

    columnas_numericas = obtener_columnas_numericas(df)
    columnas_categoricas = obtener_columnas_categoricas(df)

    if len(columnas_numericas) == 0:
        st.error("El archivo no tiene columnas numéricas para analizar.")
        st.stop()

    variable_objetivo = st.sidebar.selectbox(
        "Selecciona la variable objetivo numérica",
        columnas_numericas
    )

    st.subheader("📋 Vista general de los datos")
    st.dataframe(df, use_container_width=True)

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Registros", df.shape[0])
    col2.metric("Columnas", df.shape[1])
    col3.metric(f"Promedio de {variable_objetivo}", f"{df[variable_objetivo].mean():.2f}")
    col4.metric(f"Máximo de {variable_objetivo}", f"{df[variable_objetivo].max():.2f}")

    st.markdown("---")

    tab1, tab2 = st.tabs([
        "📆 Métodos de Pronóstico Simple",
        "📊 Regresión Lineal Simple y Múltiple"
    ])

    # ======================================================
    # TAB 1: SERIES DE TIEMPO
    # ======================================================

    with tab1:

        st.header("📆 Análisis de Series de Tiempo y Pronósticos")
        st.markdown("""
        ### 📌 ¿Qué realiza este módulo?

        Este módulo analiza una serie de datos numéricos como si fueran una **serie de tiempo**.  
        Una serie de tiempo es un conjunto de datos ordenados de manera secuencial, por ejemplo: ventas por día, rendimiento por periodo, precios por fecha o producción por mes.

        El sistema divide los datos en dos partes:

        - **Datos de entrenamiento:** sirven para construir los modelos.
        - **Datos de prueba:** sirven para verificar qué tan bien pronostica cada modelo.

        Luego, compara varios métodos de pronóstico y selecciona automáticamente el mejor modelo según el menor error **RMSE**.
        """)

        st.latex(r"Error = Valor\ real - Valor\ pronosticado")

        st.info(
            "Interpretación general: mientras menor sea el error del modelo, mejor será su capacidad "
            "para pronosticar valores futuros."
        )

        columnas_fecha = [
            col for col in df.columns
            if str(col).lower() in ["work_year", "year", "año", "anio", "fecha", "date", "tiempo", "periodo"]
        ]

        usar_columna_fecha = st.checkbox(
            "Usar columna temporal para ordenar",
            value=len(columnas_fecha) > 0
        )

        if usar_columna_fecha and columnas_fecha:
            col_tiempo = st.selectbox("Selecciona columna temporal", columnas_fecha)
            df_temp = df.sort_values(by=col_tiempo).reset_index(drop=True)
            eje_x = df_temp[col_tiempo].astype(str).values
        else:
            df_temp = df.reset_index(drop=True)
            eje_x = np.arange(len(df_temp))

        y_ts = df_temp[variable_objetivo].dropna().values.astype(float)

        if len(y_ts) < 5:
            st.warning("Se necesitan al menos 5 datos para realizar un pronóstico básico.")
        else:
            col_a, col_b, col_c = st.columns(3)

            with col_a:
                test_size = st.slider("Porcentaje de prueba", 0.1, 0.5, 0.2)

            with col_b:
                ventana = st.slider("Ventana para media móvil", 1, 10, 3)

            with col_c:
                temporada = st.slider("Temporada para ingenuo estacional", 1, 12, 3)
            st.markdown("""
            ### ⚙️ Interpretación de los parámetros

            - **Porcentaje de prueba:** indica qué parte de los datos se usará para evaluar los modelos.
            - **Ventana para media móvil:** indica cuántos datos anteriores se promedian para generar el pronóstico.
            - **Temporada:** indica cada cuántos periodos se repite un comportamiento estacional.

            Por ejemplo, si la ventana es 3, el sistema toma los últimos 3 valores para calcular el siguiente pronóstico.
            """)

            st.latex(r"\text{Entrenamiento} = 1 - \text{Porcentaje de prueba}")    

            tipo_pronostico = st.selectbox(
                "Selecciona el periodo de pronóstico",
                ["Personalizado", "Semana", "Mes", "Bimestre", "Trimestre", "Año"]
            )

            if tipo_pronostico == "Semana":
                horizonte_futuro = 7
            elif tipo_pronostico == "Mes":
                horizonte_futuro = 30
            elif tipo_pronostico == "Bimestre":
                horizonte_futuro = 60
            elif tipo_pronostico == "Trimestre":
                horizonte_futuro = 90
            elif tipo_pronostico == "Año":
                horizonte_futuro = 365
            else:
                horizonte_futuro = st.number_input("Cantidad de periodos a pronosticar", min_value=1, max_value=365, value=7)

            split = int(len(y_ts) * (1 - test_size))
            train = y_ts[:split]
            test = y_ts[split:]

            if len(test) == 0:
                st.warning("El conjunto de prueba quedó vacío. Ajusta el porcentaje.")
            else:
                horizonte_test = len(test)
                st.markdown("## 📚 Modelos de pronóstico utilizados")

                st.markdown("""
                A continuación, el sistema aplicará cinco métodos de pronóstico simple.  
                Cada método tiene una forma diferente de calcular los valores futuros.
                """)

                with st.expander("📌 Método Ingenuo"):
                    st.markdown("""
                    El **método ingenuo** utiliza el último valor observado en los datos de entrenamiento y lo repite como pronóstico.

                    Es un método simple, pero útil cuando la serie no presenta cambios fuertes o cuando se desea una comparación básica.
                    """)
                    st.latex(r"\hat{Y}_{t+1}=Y_t")
                    st.info(
                        "Interpretación: si el último valor conocido fue alto, el pronóstico también será alto. "
                        "Este método asume que el futuro inmediato será igual al último dato observado."
                    )

                with st.expander("📌 Método de la Media"):
                    st.markdown("""
                    El **método de la media** calcula el promedio de todos los datos de entrenamiento y usa ese valor como pronóstico.

                    Es útil cuando los datos no presentan una tendencia clara y se mantienen alrededor de un valor promedio.
                    """)
                    st.latex(r"\hat{Y}_{t+1}=\bar{Y}")
                    st.latex(r"\bar{Y}=\frac{Y_1+Y_2+\cdots+Y_n}{n}")
                    st.info(
                        "Interpretación: este método considera que el comportamiento futuro será parecido "
                        "al promedio histórico de la serie."
                    )

                with st.expander("📌 Método de la Media Móvil Simple"):
                    st.markdown("""
                    El **método de media móvil simple** calcula el promedio de los últimos valores de la serie, según una ventana definida.

                    A diferencia del método de la media, este modelo da más importancia a los datos recientes.
                    """)
                    st.latex(r"\hat{Y}_{t+1}=\frac{Y_t+Y_{t-1}+\cdots+Y_{t-k+1}}{k}")
                    st.info(
                        "Interpretación: si la ventana es pequeña, el modelo reacciona rápido a cambios recientes. "
                        "Si la ventana es grande, el pronóstico será más suavizado."
                    )

                with st.expander("📌 Método de Deriva"):
                    st.markdown("""
                    El **método de deriva** proyecta el comportamiento futuro usando la pendiente entre el primer y el último valor de entrenamiento.

                    Este método es útil cuando la serie muestra una tendencia creciente o decreciente.
                    """)
                    st.latex(r"\hat{Y}_{t+h}=Y_t+h\left(\frac{Y_t-Y_1}{t-1}\right)")
                    st.info(
                        "Interpretación: si los datos han crecido con el tiempo, el método de deriva continuará "
                        "esa tendencia hacia el futuro. Si han bajado, proyectará una disminución."
                    )

                with st.expander("📌 Método Ingenuo Estacional"):
                    st.markdown("""
                    El **método ingenuo estacional** utiliza valores anteriores de la misma temporada para pronosticar.

                    Se usa cuando los datos presentan patrones repetitivos cada cierto número de periodos.
                    """)
                    st.latex(r"\hat{Y}_{t+h}=Y_{t+h-m}")
                    st.info(
                        "Interpretación: este método asume que el comportamiento futuro será similar al comportamiento "
                        "observado en la misma posición de una temporada anterior."
                    )
                pronosticos_test = {
                    "Método Ingenuo": metodo_ingenuo(train, horizonte_test),
                    "Método de la Media": metodo_media(train, horizonte_test),
                    "Media Móvil Simple": metodo_media_movil(train, horizonte_test, ventana),
                    "Método de Deriva": metodo_deriva(train, horizonte_test),
                    "Ingenuo Estacional": metodo_ingenuo_estacional(train, horizonte_test, temporada)
                }

                resultados = []

                for nombre, pred in pronosticos_test.items():
                    mae = mean_absolute_error(test, pred)
                    rmse = np.sqrt(mean_squared_error(test, pred))
                    mape = mape_seguro(test, pred)

                    resultados.append({
                        "Modelo": nombre,
                        "MAE": mae,
                        "RMSE": rmse,
                        "MAPE (%)": mape
                    })

                df_resultados = pd.DataFrame(resultados).sort_values(by="RMSE")
                mejor_modelo = df_resultados.iloc[0]["Modelo"]

                st.subheader("📊 Tabla comparativa de errores")
                st.dataframe(df_resultados, use_container_width=True)
                st.markdown("### 📌 Interpretación de las métricas de error")

                st.markdown("""
                La tabla compara el rendimiento de cada método usando tres métricas principales:

                - **MAE:** mide el error promedio absoluto.
                - **RMSE:** penaliza más los errores grandes.
                - **MAPE:** expresa el error en porcentaje.

                El sistema elige como mejor modelo aquel que tenga el menor **RMSE**, porque representa un menor error general de pronóstico.
                """)

                st.latex(r"MAE=\frac{1}{n}\sum_{i=1}^{n}|Y_i-\hat{Y}_i|")

                st.latex(r"RMSE=\sqrt{\frac{1}{n}\sum_{i=1}^{n}(Y_i-\hat{Y}_i)^2}")

                st.latex(r"MAPE=\frac{100}{n}\sum_{i=1}^{n}\left|\frac{Y_i-\hat{Y}_i}{Y_i}\right|")

                st.info(
                    "Interpretación: si MAE, RMSE y MAPE son bajos, significa que el modelo se acerca bastante "
                    "a los valores reales. Si son altos, el modelo no está pronosticando correctamente."
                )
                st.success(f"✅ El mejor modelo es: {mejor_modelo}, porque tiene el menor RMSE.")
                st.markdown("### 🏆 Interpretación del mejor modelo")

                st.write(
                    f"El modelo seleccionado fue **{mejor_modelo}** porque obtuvo el menor valor de **RMSE** "
                    "en comparación con los demás métodos. Esto indica que sus pronósticos fueron los más cercanos "
                    "a los valores reales del conjunto de prueba."
                )

                if mejor_modelo == "Método Ingenuo":
                    st.info(
                        "Este resultado indica que el último valor observado representa bien el comportamiento futuro de la serie."
                    )
                elif mejor_modelo == "Método de la Media":
                    st.info(
                        "Este resultado indica que la serie se comporta alrededor de un promedio estable, sin mucha tendencia."
                    )
                elif mejor_modelo == "Media Móvil Simple":
                    st.info(
                        "Este resultado indica que los valores recientes son importantes para explicar el comportamiento futuro."
                    )
                elif mejor_modelo == "Método de Deriva":
                    st.info(
                        "Este resultado indica que la serie tiene una tendencia clara, ya sea creciente o decreciente."
                    )
                elif mejor_modelo == "Ingenuo Estacional":
                    st.info(
                        "Este resultado indica que la serie puede tener un patrón repetitivo o estacional."
                    )
                st.subheader("📈 Gráfica general de pronósticos")
                st.markdown("""
                En la siguiente gráfica se comparan los datos reales con los pronósticos generados por cada método.

                - La línea de **entrenamiento** representa los datos usados para construir los modelos.
                - La línea de **prueba real** representa los valores que el sistema intenta predecir.
                - Las líneas punteadas representan los pronósticos de cada método.

                Mientras más cerca esté la línea del pronóstico a la línea real de prueba, mejor será el modelo.
                """) 

                fig, ax = plt.subplots(figsize=(14, 6))

                x_train = np.arange(len(train))
                x_test = np.arange(len(train), len(train) + len(test))

                ax.plot(x_train, train, label="Entrenamiento", linewidth=2)
                ax.plot(x_test, test, label="Real prueba", linewidth=2)

                for nombre, pred in pronosticos_test.items():
                    ax.plot(x_test, pred, linestyle="--", label=nombre)

                ax.set_title("Comparación de métodos de pronóstico")
                ax.set_xlabel("Periodo")
                ax.set_ylabel(variable_objetivo)
                ax.legend()
                ax.grid(True, alpha=0.3)

                st.pyplot(fig)

                st.subheader("📉 Gráficas individuales por método")

                nombres_modelos = list(pronosticos_test.keys())

                for i in range(0, len(nombres_modelos), 2):
                    columnas = st.columns(2)

                    for j, col in enumerate(columnas):
                        if i + j < len(nombres_modelos):
                            nombre = nombres_modelos[i + j]
                            pred = pronosticos_test[nombre]

                            with col:
                                fig_ind, ax_ind = plt.subplots(figsize=(7, 4))
                                ax_ind.plot(x_train, train, label="Entrenamiento")
                                ax_ind.plot(x_test, test, label="Real")
                                ax_ind.plot(x_test, pred, linestyle="--", label="Pronóstico")
                                ax_ind.set_title(nombre)
                                ax_ind.set_xlabel("Periodo")
                                ax_ind.set_ylabel(variable_objetivo)
                                ax_ind.legend()
                                ax_ind.grid(True, alpha=0.3)
                                st.pyplot(fig_ind)
                                if nombre == "Método Ingenuo":
                                    st.caption(
                                        "Interpretación: este método mantiene constante el último valor conocido. "
                                        "Es útil cuando la serie es estable."
                                    )
                                elif nombre == "Método de la Media":
                                    st.caption(
                                        "Interpretación: este método usa el promedio histórico. "
                                        "Funciona mejor cuando no existe una tendencia marcada."
                                    )
                                elif nombre == "Media Móvil Simple":
                                    st.caption(
                                        "Interpretación: este método usa los últimos valores de la serie. "
                                        "Es útil cuando los datos recientes tienen mayor importancia."
                                    )
                                elif nombre == "Método de Deriva":
                                    st.caption(
                                        "Interpretación: este método proyecta la tendencia observada desde el inicio hasta el final."
                                    )
                                elif nombre == "Ingenuo Estacional":
                                    st.caption(
                                        "Interpretación: este método repite patrones anteriores según la temporada seleccionada."
                                    )

                st.subheader("🔮 Pronóstico futuro con el mejor modelo")
                st.markdown("""
                En esta sección, el sistema utiliza el mejor modelo encontrado para proyectar valores futuros.

                El horizonte de pronóstico puede ser personalizado o seleccionado como semana, mes, bimestre, trimestre o año.
                """)
                if mejor_modelo == "Método Ingenuo":
                    pronostico_futuro = metodo_ingenuo(y_ts, horizonte_futuro)
                elif mejor_modelo == "Método de la Media":
                    pronostico_futuro = metodo_media(y_ts, horizonte_futuro)
                elif mejor_modelo == "Media Móvil Simple":
                    pronostico_futuro = metodo_media_movil(y_ts, horizonte_futuro, ventana)
                elif mejor_modelo == "Método de Deriva":
                    pronostico_futuro = metodo_deriva(y_ts, horizonte_futuro)
                else:
                    pronostico_futuro = metodo_ingenuo_estacional(y_ts, horizonte_futuro, temporada)

                df_futuro = pd.DataFrame({
                    "Periodo futuro": np.arange(1, horizonte_futuro + 1),
                    "Pronóstico": pronostico_futuro
                })

                st.dataframe(df_futuro, use_container_width=True)
                st.markdown("### 📌 Interpretación del pronóstico futuro")

                st.write(
                    f"El sistema generó un pronóstico de **{horizonte_futuro} periodos futuros** usando el modelo "
                    f"**{mejor_modelo}**. Estos valores representan una estimación del comportamiento esperado "
                    f"de la variable **{variable_objetivo}**."
                )

                st.warning(
                    "Importante: un pronóstico no es un valor exacto, sino una estimación basada en el comportamiento histórico "
                    "de los datos. Por eso debe interpretarse como apoyo para la toma de decisiones."
                )
                st.info(
                    f"Interpretación: El sistema eligió el modelo {mejor_modelo} "
                    f"porque obtuvo el menor error RMSE. Se generó un pronóstico de "
                    f"{horizonte_futuro} periodos futuros."
                )

    # ======================================================
    # TAB 2: REGRESIÓN
    # ======================================================

    with tab2:

        st.header("📊 Regresión Lineal Simple y Múltiple")

        df_reg = df.copy()
        df_reg = df_reg.dropna(subset=[variable_objetivo])

        columnas_numericas_reg = obtener_columnas_numericas(df_reg)
        columnas_categoricas_reg = obtener_columnas_categoricas(df_reg)

        columnas_predictoras_numericas = [c for c in columnas_numericas_reg if c != variable_objetivo]

        # --------------------------------------------------
        # REGRESIÓN LINEAL SIMPLE
        # --------------------------------------------------

        st.subheader("1️⃣ Regresión Lineal Simple")

        if len(columnas_predictoras_numericas) == 0:
            st.warning("No hay variables numéricas predictoras para regresión simple.")
        else:
            correlaciones = df_reg[columnas_predictoras_numericas + [variable_objetivo]].corr(numeric_only=True)[variable_objetivo]
            correlaciones = correlaciones.drop(variable_objetivo).abs().sort_values(ascending=False)

            mejor_x = correlaciones.index[0]

            st.write(f"Variable seleccionada automáticamente por mayor correlación: **{mejor_x}**")

            data_simple = df_reg[[mejor_x, variable_objetivo]].dropna()

            X_simple = data_simple[[mejor_x]]
            y_simple = data_simple[variable_objetivo]

            modelo_simple = LinearRegression()
            modelo_simple.fit(X_simple, y_simple)

            y_pred_simple = modelo_simple.predict(X_simple)

            a_simple = modelo_simple.intercept_
            b_simple = modelo_simple.coef_[0]
            r2_simple = r2_score(y_simple, y_pred_simple)

            col1, col2, col3 = st.columns(3)
            col1.metric("Intercepto a", f"{a_simple:.4f}")
            col2.metric("Coeficiente b", f"{b_simple:.4f}")
            col3.metric("R²", f"{r2_simple:.4f}")

            st.code(f"{variable_objetivo} = {a_simple:.4f} + ({b_simple:.4f}) * {mejor_x}")

            fig_simple, ax_simple = plt.subplots(figsize=(10, 5))
            ax_simple.scatter(X_simple, y_simple, alpha=0.7)
            ax_simple.plot(X_simple, y_pred_simple, linewidth=2)
            ax_simple.set_title("Regresión lineal simple")
            ax_simple.set_xlabel(mejor_x)
            ax_simple.set_ylabel(variable_objetivo)
            ax_simple.grid(True, alpha=0.3)
            st.pyplot(fig_simple)

            try:
                X_sm = sm.add_constant(X_simple)
                modelo_sm = sm.OLS(y_simple, X_sm).fit()

                # ===============================
                # TABLA ANOVA MANUAL Y DINÁMICA
                # ===============================

                y_real = np.array(y_simple)
                y_estimado = np.array(y_pred_simple)
                media_y = np.mean(y_real)

                # Sumas de cuadrados
                SSR = np.sum((y_estimado - media_y) ** 2)   # Regresión
                SSE = np.sum((y_real - y_estimado) ** 2)    # Error
                SST = np.sum((y_real - media_y) ** 2)       # Total

                # Grados de libertad
                gl_regresion = 1
                gl_error = len(y_real) - 2
                gl_total = len(y_real) - 1

                # Cuadrados medios
                MSR = SSR / gl_regresion if gl_regresion != 0 else np.nan
                MSE = SSE / gl_error if gl_error != 0 else np.nan

                # Estadístico F y p-valor
                F = MSR / MSE if MSE != 0 else np.nan
                p_valor_anova = 1 - stats.f.cdf(F, gl_regresion, gl_error) if not np.isnan(F) else np.nan

                tabla_anova = pd.DataFrame({
                    "Fuente": ["Regresión", "Error", "Total"],
                    "Suma de cuadrados": [SSR, SSE, SST],
                    "Grados de libertad": [gl_regresion, gl_error, gl_total],
                    "Cuadrado medio": [MSR, MSE, ""],
                    "F": [F, "", ""],
                    "p-valor": [p_valor_anova, "", ""]
                })

                st.subheader("📌 Tabla ANOVA")
                st.dataframe(tabla_anova, use_container_width=True)

                st.markdown("""
                **Interpretación de la tabla ANOVA:**  
                La tabla ANOVA permite evaluar si el modelo de regresión lineal simple es estadísticamente significativo.

                - **Regresión:** representa la variabilidad explicada por el modelo.
                - **Error:** representa la variabilidad que el modelo no logra explicar.
                - **Total:** representa la variabilidad total de los datos.
                - **F:** permite comparar la variabilidad explicada frente al error.
                - **p-valor:** si es menor a 0.05, se considera que el modelo es significativo.
                """)

                if p_valor_anova < 0.05:
                    st.success(
                        "Como el p-valor es menor a 0.05, se acepta que el modelo tiene significancia estadística."
                    )
                else:
                    st.warning(
                        "Como el p-valor es mayor o igual a 0.05, el modelo no presenta suficiente significancia estadística."
                    )

                # ===============================
                # P-VALORES DE LOS COEFICIENTES
                # ===============================

                st.subheader("📌 P-valores de los coeficientes")

                tabla_pvalores = pd.DataFrame({
                    "Coeficiente": modelo_sm.params.index,
                    "Valor": modelo_sm.params.values,
                    "p-valor": modelo_sm.pvalues.values
                })

                st.dataframe(tabla_pvalores, use_container_width=True)

                st.markdown("""
                **Interpretación de los p-valores:**  
                Los p-valores permiten analizar si cada coeficiente aporta significativamente al modelo.

                - Si el p-valor de una variable es menor a 0.05, esa variable influye significativamente.
                - Si el p-valor es mayor a 0.05, su influencia no es tan fuerte estadísticamente.
                """)

            except Exception as e:
                st.warning(f"No se pudo generar ANOVA: {e}")

            st.info(
                f"Interpretación: La variable {mejor_x} fue seleccionada porque tiene la mayor relación "
                f"con {variable_objetivo}. El valor R² indica qué proporción de variabilidad explica el modelo."
            )

            st.markdown("---")
        # --------------------------------------------------
        # REGRESIÓN LINEAL MÚLTIPLE
        # --------------------------------------------------

        st.subheader("2️⃣ Regresión Lineal Múltiple")

        columnas_disponibles = [c for c in df_reg.columns if c != variable_objetivo]

        predictoras = st.multiselect(
            "Selecciona variables predictoras",
            columnas_disponibles,
            default=columnas_disponibles[:min(4, len(columnas_disponibles))]
        )

        r2_multiple = None
        mae_multiple = None
        accuracy_clasificacion = None

        if len(predictoras) == 0:
            st.warning("Selecciona al menos una variable predictora.")
        else:
            data_modelo = df_reg[predictoras + [variable_objetivo]].dropna()

            if len(data_modelo) < 5:
                st.warning("Se necesitan al menos 5 registros para entrenar el modelo.")
            else:
                X = data_modelo[predictoras]
                y = data_modelo[variable_objetivo]

                num_features = X.select_dtypes(include=[np.number]).columns.tolist()
                cat_features = X.select_dtypes(include=["object", "category", "bool"]).columns.tolist()

                preprocesador = ColumnTransformer(
                    transformers=[
                        ("num", "passthrough", num_features),
                        ("cat", OneHotEncoder(handle_unknown="ignore", sparse_output=False), cat_features)
                    ],
                    remainder="drop"
                )

                modelo_multiple = Pipeline(
                    steps=[
                        ("preprocesador", preprocesador),
                        ("modelo", LinearRegression())
                    ]
                )

                X_train, X_test, y_train, y_test = train_test_split(
                    X, y, test_size=0.25, random_state=42
                )

                modelo_multiple.fit(X_train, y_train)
                y_pred = modelo_multiple.predict(X_test)

                mae_multiple, mse_multiple, rmse_multiple, r2_multiple = calcular_metricas_regresion(y_test, y_pred)

                col1, col2, col3, col4 = st.columns(4)
                col1.metric("R²", f"{r2_multiple:.4f}")
                col2.metric("MAE", f"{mae_multiple:.4f}")
                col3.metric("MSE", f"{mse_multiple:.4f}")
                col4.metric("RMSE", f"{rmse_multiple:.4f}")

                st.subheader("📌 Ecuación del modelo múltiple")

                try:
                    nombres_features = modelo_multiple.named_steps["preprocesador"].get_feature_names_out()
                    coeficientes = modelo_multiple.named_steps["modelo"].coef_
                    intercepto = modelo_multiple.named_steps["modelo"].intercept_

                    ecuacion = f"{variable_objetivo} = {intercepto:.4f}"
                    tabla_coef = []

                    for nombre, coef in zip(nombres_features, coeficientes):
                        nombre_limpio = nombre.replace("num__", "").replace("cat__", "")
                        ecuacion += f" + ({coef:.4f})*{nombre_limpio}"
                        tabla_coef.append({
                            "Variable": nombre_limpio,
                            "Coeficiente": coef
                        })

                    st.code(ecuacion)
                    st.dataframe(pd.DataFrame(tabla_coef), use_container_width=True)

                except Exception as e:
                    st.warning(f"No se pudo mostrar ecuación completa: {e}")

                st.info(
                    "Interpretación: Cada coeficiente indica cuánto cambia la variable objetivo "
                    "cuando la variable predictora aumenta una unidad, manteniendo las demás constantes."
                )

                # --------------------------------------------------
                # PREDICCIÓN CON NUEVOS DATOS
                # --------------------------------------------------

                st.subheader("🔮 Predicción con nuevos datos")

                with st.form("form_prediccion"):
                    nuevos_datos = {}

                    for col in predictoras:
                        if col in num_features:
                            valor_defecto = float(df_reg[col].mean())
                            nuevos_datos[col] = st.number_input(f"Ingrese {col}", value=valor_defecto)
                        else:
                            opciones = df_reg[col].dropna().astype(str).unique().tolist()
                            if len(opciones) == 0:
                                opciones = ["Sin dato"]
                            nuevos_datos[col] = st.selectbox(f"Seleccione {col}", opciones)

                    boton_pred = st.form_submit_button("Predecir")

                if boton_pred:
                    df_nuevo = pd.DataFrame([nuevos_datos])
                    pred_nuevo = modelo_multiple.predict(df_nuevo)[0]
                    st.success(f"Predicción estimada de {variable_objetivo}: {pred_nuevo:.4f}")

                # --------------------------------------------------
                # VIF
                # --------------------------------------------------

                st.subheader("📌 Diagnóstico de multicolinealidad VIF")

                if len(num_features) < 2:
                    st.info("Se necesitan al menos dos variables numéricas para calcular VIF.")
                else:
                    try:
                        X_vif = data_modelo[num_features].dropna()
                        X_vif_const = sm.add_constant(X_vif)

                        vif_data = pd.DataFrame()
                        vif_data["Variable"] = X_vif_const.columns
                        vif_data["VIF"] = [
                            variance_inflation_factor(X_vif_const.values, i)
                            for i in range(X_vif_const.shape[1])
                        ]

                        st.dataframe(vif_data, use_container_width=True)

                        variables_altas = vif_data[(vif_data["Variable"] != "const") & (vif_data["VIF"] > 10)]["Variable"].tolist()

                        if variables_altas:
                            st.warning(
                                f"Variables con VIF mayor a 10: {variables_altas}. "
                                "Se recomienda revisar o eliminar estas variables."
                            )
                        else:
                            st.success("No se detecta multicolinealidad fuerte según VIF.")

                    except Exception as e:
                        st.warning(f"No se pudo calcular VIF: {e}")

                # --------------------------------------------------
                # RESIDUOS
                # --------------------------------------------------

                st.subheader("📌 Diagnóstico de residuos")

                residuos = y_test - y_pred

                col_res1, col_res2 = st.columns(2)

                with col_res1:
                    fig_qq = plt.figure(figsize=(6, 4))
                    sm.qqplot(residuos, line="45", fit=True)
                    plt.title("Q-Q Plot de residuos")
                    st.pyplot(fig_qq)

                with col_res2:
                    fig_res, ax_res = plt.subplots(figsize=(6, 4))
                    ax_res.scatter(y_pred, residuos, alpha=0.7)
                    ax_res.axhline(0, linestyle="--")
                    ax_res.set_xlabel("Valores predichos")
                    ax_res.set_ylabel("Residuos")
                    ax_res.set_title("Residuos vs predichos")
                    ax_res.grid(True, alpha=0.3)
                    st.pyplot(fig_res)

                # --------------------------------------------------
                # GRÁFICOS ADICIONALES
                # --------------------------------------------------

                st.subheader("📊 Gráficos adicionales para presentación")

                if len(columnas_numericas_reg) >= 2:
                    try:
                        cols_pair = columnas_numericas_reg[:5]
                        fig_pair = sns.pairplot(df_reg[cols_pair].dropna())
                        st.pyplot(fig_pair.fig)
                    except Exception as e:
                        st.warning(f"No se pudo generar pairplot: {e}")
                else:
                    st.info("No hay suficientes variables numéricas para pairplot.")

                if len(columnas_categoricas_reg) > 0:
                    cat_box = st.selectbox("Selecciona variable categórica para boxplot", columnas_categoricas_reg)

                    try:
                        fig_box, ax_box = plt.subplots(figsize=(10, 5))
                        sns.boxplot(data=df_reg, x=cat_box, y=variable_objetivo, ax=ax_box)
                        ax_box.set_title(f"Boxplot de {variable_objetivo} según {cat_box}")
                        ax_box.tick_params(axis="x", rotation=45)
                        st.pyplot(fig_box)
                    except Exception as e:
                        st.warning(f"No se pudo generar boxplot: {e}")
                else:
                    st.info("No hay variables categóricas para boxplot.")

                if "salary" in df_reg.columns and "salary_in_usd" in df_reg.columns:
                    try:
                        fig_salary, ax_salary = plt.subplots(figsize=(8, 5))
                        ax_salary.scatter(df_reg["salary"], df_reg["salary_in_usd"], alpha=0.7)
                        ax_salary.set_xlabel("salary")
                        ax_salary.set_ylabel("salary_in_usd")
                        ax_salary.set_title("Relación salary vs salary_in_usd")
                        ax_salary.grid(True, alpha=0.3)
                        st.pyplot(fig_salary)
                    except:
                        pass

                # --------------------------------------------------
                # MATRIZ DE CONFUSIÓN
                # --------------------------------------------------

                st.subheader("🧮 Matriz de Confusión y Métricas de Clasificación")

                try:
                    q33 = np.quantile(y_test, 0.33)
                    q66 = np.quantile(y_test, 0.66)

                    def categorizar(valor):
                        if valor <= q33:
                            return "Bajo"
                        elif valor <= q66:
                            return "Medio"
                        else:
                            return "Alto"

                    y_test_cat = np.array([categorizar(v) for v in y_test])
                    y_pred_cat = np.array([categorizar(v) for v in y_pred])

                    labels = ["Bajo", "Medio", "Alto"]

                    cm = confusion_matrix(y_test_cat, y_pred_cat, labels=labels)

                    accuracy_clasificacion = accuracy_score(y_test_cat, y_pred_cat)
                    precision = precision_score(y_test_cat, y_pred_cat, average="weighted", zero_division=0)
                    recall = recall_score(y_test_cat, y_pred_cat, average="weighted", zero_division=0)
                    f1 = f1_score(y_test_cat, y_pred_cat, average="weighted", zero_division=0)

                    col1, col2, col3, col4 = st.columns(4)
                    col1.metric("Exactitud", f"{accuracy_clasificacion:.4f}")
                    col2.metric("Precisión", f"{precision:.4f}")
                    col3.metric("Sensibilidad", f"{recall:.4f}")
                    col4.metric("F1-score", f"{f1:.4f}")

                    fig_cm, ax_cm = plt.subplots(figsize=(7, 5))
                    sns.heatmap(cm, annot=True, fmt="d", cmap="Blues", xticklabels=labels, yticklabels=labels, ax=ax_cm)
                    ax_cm.set_xlabel("Predicción")
                    ax_cm.set_ylabel("Real")
                    ax_cm.set_title("Matriz de Confusión")
                    st.pyplot(fig_cm)

                    st.info(
                        "Regla aplicada: La variable objetivo numérica fue convertida en categorías "
                        "Bajo, Medio y Alto usando los cuantiles 33% y 66%. Por eso se puede aplicar "
                        "matriz de confusión."
                    )

                    # Especificidad aproximada por clase
                    st.subheader("📌 Especificidad por categoría")

                    especificidades = []
                    total = cm.sum()

                    for i, label in enumerate(labels):
                        TP = cm[i, i]
                        FP = cm[:, i].sum() - TP
                        FN = cm[i, :].sum() - TP
                        TN = total - TP - FP - FN
                        especificidad = TN / (TN + FP) if (TN + FP) != 0 else 0

                        especificidades.append({
                            "Categoría": label,
                            "Especificidad": especificidad
                        })

                    st.dataframe(pd.DataFrame(especificidades), use_container_width=True)

                except Exception as e:
                    st.warning(f"No se pudo generar matriz de confusión: {e}")

                # --------------------------------------------------
                # REPORTE PDF
                # --------------------------------------------------

                st.subheader("📄 Reporte PDF")

                pdf_buffer = crear_pdf_reporte(
                    variable_objetivo,
                    r2_multiple,
                    mae_multiple,
                    accuracy_clasificacion
                )

                st.download_button(
                    label="📥 Descargar reporte PDF",
                    data=pdf_buffer,
                    file_name="reporte_sistema_predictivo.pdf",
                    mime="application/pdf"
                )

# ==========================================================
# FUNCIONALIDAD 2: DOCUMENTOS DEL INGE
# ==========================================================

elif opcion_principal == "📄 Documentos del Inge":

    st.header("📄 Funcionalidad 2: Análisis de Documentos del Inge")

    documento = st.file_uploader(
        "Sube un documento PDF, Word o PowerPoint",
        type=["pdf", "docx", "pptx"]
    )

    if documento is None:
        st.info("Sube un documento del docente para analizarlo.")
        st.stop()

    texto = ""

    try:
        if documento.name.endswith(".pdf"):
            texto = leer_pdf(documento)
        elif documento.name.endswith(".docx"):
            texto = leer_docx(documento)
        elif documento.name.endswith(".pptx"):
            texto = leer_pptx(documento)
        else:
            st.error("Formato no compatible.")
            st.stop()

    except Exception as e:
        st.error(f"No se pudo leer el documento: {e}")
        st.stop()

    st.success("Documento leído correctamente")

    temas = detectar_tema(texto)
    formulas = extraer_posibles_formulas(texto)

    col1, col2, col3 = st.columns(3)
    col1.metric("Caracteres extraídos", len(texto))
    col2.metric("Temas detectados", len(temas))
    col3.metric("Fórmulas/patrones", len(formulas))

    st.subheader("📌 Temas detectados")
    for tema in temas:
        st.write(f"✅ {tema}")

    st.subheader("🧾 Resumen automático básico")

    parrafos = [p.strip() for p in texto.split("\n") if len(p.strip()) > 50]
    resumen = " ".join(parrafos[:5])

    if resumen:
        st.write(resumen[:2000])
    else:
        st.info("No se pudo generar resumen porque el documento tiene poco texto extraíble.")

    st.subheader("📐 Fórmulas o patrones encontrados")

    if formulas:
        for f in formulas:
            st.code(f)
    else:
        st.info("No se detectaron fórmulas automáticamente.")

    st.subheader("📖 Texto extraído del documento")

    with st.expander("Ver texto completo"):
        st.text_area("Contenido", texto, height=500)

    st.info(
        "Interpretación: Este módulo permite cargar documentos del docente, extraer su contenido, "
        "detectar temas principales y ubicar fórmulas o conceptos importantes para el desarrollo del sistema."
    )
